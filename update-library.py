"""
BGG Library Updater

Run this script whenever you add new games to your BGG collection:
    python update-library.py

The script will:
    1. Fetch your current collection from BoardGameGeek
    2. Fetch detailed info for each game (description, designers, etc.)
    3. Pull in our own shelf info (genre, complexity, small area, 18+) from
       the "Shelf Info" tab of game_library_master.xlsx, and add any new
       games to that tab
    4. Save everything to games.json and push it to GitHub

After filling in the spreadsheet, publish just those changes (no BGG fetch):
    python update-library.py --sheet-only

Add --no-push to either command to update games.json without publishing it.
"""

import functools
import json
import os
import re
import html
import sys
import time
import subprocess
import urllib.request
import urllib.error
from datetime import datetime

# ============ CONFIGURATION ============
BGG_USERNAME = 'cavernsandforests'
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, 'games.json')
# The API token is NOT stored in this file (it's published to GitHub).
# Set the BGG_API_TOKEN environment variable, or keep the token in
# bgg_xml_api_access_token.txt next to this script (git-ignored).
TOKEN_FILE = os.path.join(SCRIPT_DIR, 'bgg_xml_api_access_token.txt')

# Our own shelf info is typed into this spreadsheet. It stays on this
# computer; only games.json is published.
SHEET_FILE = os.path.join(SCRIPT_DIR, 'game_library_master.xlsx')
SHEET_NAME = 'Shelf Info'
GENRES_SHEET = 'Genres'
COMPLEXITY_LEVELS = ['Light', 'Medium', 'Challenging', 'Epic']
# Custom fields we fill in ourselves (BGG never overwrites them), with defaults
CUSTOM_FIELDS = {
    'ourComplexity': '',
    'shelfLocation': '',  # our genre: where the game lives on the shelf
    'smallArea': False,
    'adultsOnly': False,  # 18+
}
# (header, column width, typed by us?) for the Shelf Info tab
SHEET_COLUMNS = [
    ('BGG ID', 10, False),
    ('Game', 46, False),
    ('Genre', 24, True),
    ('Complexity', 13, True),
    ('Small Area', 11, True),
    ('18+', 7, True),
    ('Notes', 32, True),
    ('Expansion Of', 36, False),
    ('Players', 9, False),
    ('BGG Age', 9, False),
    ('BGG Weight', 11, False),
    ('Status', 22, False),
]
# =======================================


@functools.lru_cache(maxsize=None)
def load_api_token():
    """Read the BGG API token from the environment or the local token file."""
    token = os.environ.get('BGG_API_TOKEN', '').strip()
    if token:
        return token
    if os.path.exists(TOKEN_FILE):
        with open(TOKEN_FILE, 'r', encoding='utf-8') as f:
            match = re.search(r'[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}', f.read())
            if match:
                return match.group(0)
    raise SystemExit(
        'No BGG API token found. Set BGG_API_TOKEN or put the token in '
        'bgg_xml_api_access_token.txt next to this script.'
    )


def fetch_url(url, max_retries=10):
    """Fetch a URL with retry logic for BGG's processing queue."""
    for attempt in range(max_retries):
        print(f'  Fetching... (attempt {attempt + 1}/{max_retries})')
        try:
            req = urllib.request.Request(url, headers={
                'User-Agent': 'CavernsAndForestsLibraryUpdater/1.0',
                'Authorization': f'Bearer {load_api_token()}',
            })
            with urllib.request.urlopen(req) as response:
                data = response.read().decode('utf-8')

                if 'will be processed' in data:
                    print('  BGG is processing request, waiting...')
                    time.sleep(3)
                    continue

                return data
        except urllib.error.HTTPError as e:
            if e.code == 202:
                print('  BGG is processing request, waiting...')
                time.sleep(3)
                continue
            raise Exception(f'HTTP {e.code}')
        except urllib.error.URLError as e:
            raise Exception(f'Network error: {e.reason}')

    raise Exception('Max retries exceeded - BGG took too long to respond')


def extract_links(xml_text, link_type):
    """Extract link values of a given type from BGG XML."""
    results = []
    pattern = re.compile(
        rf'link[^>]*type="{link_type}"[^>]*value="([^"]+)"', re.IGNORECASE
    )
    for match in pattern.finditer(xml_text):
        results.append(html.unescape(match.group(1)))
    return results


def parse_xml_value(xml_text, tag):
    """Extract text content from a simple XML tag."""
    match = re.search(rf'<{tag}[^>]*>([^<]*)</{tag}>', xml_text)
    return match.group(1).strip() if match else ''


def parse_xml_attr(xml_text, tag, attr):
    """Extract an attribute value from an XML tag."""
    match = re.search(rf'<{tag}[^>]*{attr}="([^"]*)"', xml_text, re.IGNORECASE)
    return match.group(1) if match else ''


def parse_collection(xml_text):
    """Parse BGG collection XML into a list of games."""
    games = []
    item_pattern = re.compile(
        r'<item[^>]*objectid="(\d+)"[^>]*>(.*?)</item>', re.DOTALL
    )

    for match in item_pattern.finditer(xml_text):
        game_id = match.group(1)
        item_xml = match.group(2)

        # Get BGG rating
        rating = 0.0
        avg_match = re.search(r'<average[^>]*value="([^"]+)"', item_xml)
        if avg_match:
            try:
                rating = float(avg_match.group(1))
            except ValueError:
                pass

        game = {
            'id': game_id,
            'name': parse_xml_value(item_xml, 'name'),
            'yearPublished': parse_xml_value(item_xml, 'yearpublished') or 'N/A',
            'image': parse_xml_value(item_xml, 'image'),
            'thumbnail': parse_xml_value(item_xml, 'thumbnail'),
            'minPlayers': int(parse_xml_attr(item_xml, 'stats', 'minplayers') or 1),
            'maxPlayers': int(parse_xml_attr(item_xml, 'stats', 'maxplayers') or 1),
            'playingTime': int(parse_xml_attr(item_xml, 'stats', 'playingtime') or 0),
            'rating': rating,
            # Fields populated by fetch_game_details
            'type': 'boardgame',
            'description': '',
            'shortDescription': '',
            'designers': [],
            'artists': [],
            'publishers': [],
            'categories': [],
            'mechanics': [],
            'honors': [],
            'minAge': 0,
            'bggWeight': 0.0,
            'bggRank': 0,
            'languageDependence': '',
            'baseGame': None,
            # Custom fields - preserved across updates
            **CUSTOM_FIELDS,
        }

        games.append(game)

    return games


def fetch_game_details(games):
    """Fetch detailed info for each game from the BGG Thing API."""
    batch_size = 20

    for i in range(0, len(games), batch_size):
        batch = games[i:i + batch_size]
        ids = ','.join(g['id'] for g in batch)
        end = min(i + batch_size, len(games))

        print(f'Fetching details for games {i + 1}-{end} of {len(games)}...')

        try:
            url = f'https://boardgamegeek.com/xmlapi2/thing?id={ids}&stats=1'
            xml_text = fetch_url(url)

            # Parse each item from the response
            item_pattern = re.compile(
                r'<item[^>]*id="(\d+)"[^>]*>(.*?)</item>', re.DOTALL
            )

            for match in item_pattern.finditer(xml_text):
                item_open_end = xml_text.index('>', match.start()) + 1
                item_open_tag = xml_text[match.start():item_open_end]
                game_id = match.group(1)
                item_xml = match.group(2)

                game = next((g for g in batch if g['id'] == game_id), None)
                if not game:
                    continue

                # Type (from item tag attribute)
                type_match = re.search(r'type="([^"]+)"', item_open_tag)
                game['type'] = type_match.group(1) if type_match else 'boardgame'

                # Description
                desc_match = re.search(r'<description>(.*?)</description>', item_xml, re.DOTALL)
                if desc_match:
                    # BGG double-encodes HTML entities in descriptions
                    game['description'] = html.unescape(html.unescape(desc_match.group(1)))
                    clean = re.sub(r'\s+', ' ', game['description']).strip()
                    if len(clean) > 200:
                        cutoff = clean.rfind(' ', 0, 200)
                        game['shortDescription'] = clean[:cutoff if cutoff > 0 else 200] + '...'
                    else:
                        game['shortDescription'] = clean

                # Designers, Artists, Publishers
                game['designers'] = extract_links(item_xml, 'boardgamedesigner')
                game['artists'] = extract_links(item_xml, 'boardgameartist')
                game['publishers'] = extract_links(item_xml, 'boardgamepublisher')

                # Categories & Mechanics
                game['categories'] = extract_links(item_xml, 'boardgamecategory')
                game['mechanics'] = extract_links(item_xml, 'boardgamemechanic')

                # Awards & Honors
                game['honors'] = extract_links(item_xml, 'boardgamehonor')

                # Min Age
                age_match = re.search(r'<minage[^>]*value="([^"]+)"', item_xml)
                game['minAge'] = int(age_match.group(1)) if age_match else 0

                # BGG Weight/Complexity (1-5 scale)
                weight_match = re.search(r'averageweight[^>]*value="([^"]+)"', item_xml)
                game['bggWeight'] = float(weight_match.group(1)) if weight_match else 0.0

                # BGG Rank
                rank_match = re.search(
                    r'rank[^>]*type="subtype"[^>]*name="boardgame"[^>]*value="([^"]+)"',
                    item_xml
                )
                if rank_match and rank_match.group(1) != 'Not Ranked':
                    try:
                        game['bggRank'] = int(rank_match.group(1))
                    except ValueError:
                        game['bggRank'] = 0
                else:
                    game['bggRank'] = 0

                # Language Dependence (winning vote from poll)
                lang_poll_match = re.search(
                    r'<poll[^>]*name="language_dependence"[^>]*>(.*?)</poll>',
                    item_xml, re.DOTALL
                )
                if lang_poll_match:
                    votes = []
                    result_pattern = re.compile(
                        r'<result[^>]*value="([^"]+)"[^>]*numvotes="(\d+)"'
                    )
                    for r_match in result_pattern.finditer(lang_poll_match.group(1)):
                        votes.append({
                            'value': r_match.group(1),
                            'count': int(r_match.group(2)),
                        })
                    if votes:
                        votes.sort(key=lambda v: v['count'], reverse=True)
                        game['languageDependence'] = votes[0]['value'] if votes[0]['count'] > 0 else ''

                # Base game (for expansions)
                if game['type'] == 'boardgameexpansion':
                    base_match = re.search(
                        r'<link[^>]*type="boardgameexpansion"[^>]*id="(\d+)"[^>]*value="([^"]+)"[^>]*inbound="true"',
                        item_xml
                    )
                    if base_match:
                        game['baseGame'] = {
                            'id': base_match.group(1),
                            'name': html.unescape(base_match.group(2)),
                        }

            # Small delay between batches
            if i + batch_size < len(games):
                time.sleep(1)

        except Exception as e:
            print(f'  Warning: Could not fetch details for batch: {e}')


def is_yes(value):
    """True for a ticked spreadsheet cell: Yes, Y, X, TRUE or 1."""
    return value is True or str(value or '').strip().lower() in ('yes', 'y', 'x', 'true', '1')


def sheet_id(value):
    """BGG IDs come back from Excel as numbers; normalise them to strings."""
    if isinstance(value, float):
        value = int(value)
    return str(value or '').strip()


def sync_shelf_sheet(games):
    """Sync our shelf info with the Shelf Info tab of the master spreadsheet.

    Genre, Complexity, Small Area, 18+ and Notes are typed into the
    spreadsheet, which is the source of truth for them. The gray columns are
    refreshed from BGG every run, and new games get a row of their own.
    Returns False if the spreadsheet couldn't be used.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print('Skipping the spreadsheet: run "python -m pip install openpyxl" to enable it.\n')
        return False

    if os.path.exists(SHEET_FILE):
        wb = openpyxl.load_workbook(SHEET_FILE)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    # What's been typed in so far, keyed by BGG ID
    typed = {}
    if SHEET_NAME in wb.sheetnames:
        rows = wb[SHEET_NAME].iter_rows(values_only=True)
        header = [str(h or '').strip() for h in next(rows, [])]
        for row in rows:
            rec = dict(zip(header, row))
            if sheet_id(rec.get('BGG ID')):
                typed[sheet_id(rec.get('BGG ID'))] = rec

    if GENRES_SHEET not in wb.sheetnames:
        genre_ws = wb.create_sheet(GENRES_SHEET)
        genre_ws.append(['Genre'])
        for name in sorted({g['shelfLocation'] for g in games if g.get('shelfLocation')}):
            genre_ws.append([name])
    genres = {str(r[0]).strip().lower(): str(r[0]).strip()
              for r in wb[GENRES_SHEET].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]}
    levels = {c.lower(): c for c in COMPLEXITY_LEVELS}

    # Apply the typed values. New games keep whatever games.json already had.
    # A genre or complexity that isn't on its list stays in the spreadsheet
    # but isn't published, so a typo never shows up on the website.
    problems, as_typed = {}, {}
    for game in games:
        rec = typed.get(game['id'])
        if rec is None:
            continue
        genre = str(rec.get('Genre') or '').strip()
        if genre and genre.lower() not in genres:
            problems[f'{game["name"]}: genre "{genre}" is not on the {GENRES_SHEET} tab '
                     f'(not published until fixed)'] = True
        complexity = str(rec.get('Complexity') or '').strip()
        if complexity and complexity.lower() not in levels:
            problems[f'{game["name"]}: complexity "{complexity}" should be one of '
                     f'{", ".join(COMPLEXITY_LEVELS)} (not published until fixed)'] = True
        game['shelfLocation'] = genres.get(genre.lower(), '')
        game['ourComplexity'] = levels.get(complexity.lower(), '')
        game['smallArea'] = is_yes(rec.get('Small Area'))
        game['adultsOnly'] = is_yes(rec.get('18+'))
        as_typed[game['id']] = (genres.get(genre.lower(), genre), levels.get(complexity.lower(), complexity))

    # Rebuild the tab: one row per game (owning two copies still means one row)
    rows, seen = [], set()
    for game in games:
        if game['id'] in seen:
            continue
        seen.add(game['id'])
        rec = typed.get(game['id'], {})
        if game['minPlayers'] == game['maxPlayers']:
            players = game['minPlayers']
        else:
            players = f'{game["minPlayers"]}-{game["maxPlayers"]}'
        genre, complexity = as_typed.get(game['id'], (game['shelfLocation'], game['ourComplexity']))
        rows.append([
            int(game['id']), html.unescape(game['name']),
            genre, complexity,
            'Yes' if game['smallArea'] else '', 'Yes' if game['adultsOnly'] else '',
            rec.get('Notes') or '',
            html.unescape((game.get('baseGame') or {}).get('name', '')),
            players, game.get('minAge') or '', round(game.get('bggWeight') or 0, 2) or '',
            'New' if typed and game['id'] not in typed else '',
        ])
    # Keep rows for games that left the BGG collection, so nothing typed is lost
    for game_id, rec in typed.items():
        if game_id not in seen:
            rows.append([int(game_id) if game_id.isdigit() else game_id] +
                        [rec.get(name) or '' for name, _, _ in SHEET_COLUMNS[1:-1]] +
                        ['Not in BGG collection'])
    rows.sort(key=lambda r: str(r[1]).lower())

    index = wb.sheetnames.index(SHEET_NAME) if SHEET_NAME in wb.sheetnames else 0
    if SHEET_NAME in wb.sheetnames:
        wb.remove(wb[SHEET_NAME])
    ws = wb.create_sheet(SHEET_NAME, index)
    ws.append([name for name, _, _ in SHEET_COLUMNS])
    for row in rows:
        ws.append(row)

    # Green headers = ours to fill in; gray columns come from BGG
    gray = PatternFill('solid', fgColor='EEEEEE')
    for col, (name, width, ours) in enumerate(SHEET_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        ws.cell(1, col).font = Font(bold=True, color='FFFFFF')
        ws.cell(1, col).fill = PatternFill('solid', fgColor='2F6B3A' if ours else '555555')
        if not ours:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).fill = gray
    last = ws.max_row
    col_of = {name: get_column_letter(i) for i, (name, _, _) in enumerate(SHEET_COLUMNS, 1)}
    for name, formula, message in [
        ('Genre', f'{GENRES_SHEET}!$A$2:$A$200',
         f'Pick a genre from the list (add new ones on the {GENRES_SHEET} tab).'),
        ('Complexity', '"' + ','.join(COMPLEXITY_LEVELS) + '"', 'Pick ' + ', '.join(COMPLEXITY_LEVELS) + '.'),
        ('Small Area', '"Yes"', 'Pick Yes, or leave it blank.'),
        ('18+', '"Yes"', 'Pick Yes, or leave it blank.'),
    ]:
        dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorStyle='warning', errorTitle='Not on the list', error=message)
        dv.add(f'{col_of[name]}2:{col_of[name]}{last}')
        ws.add_data_validation(dv)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(SHEET_COLUMNS))}{last}'

    try:
        wb.save(SHEET_FILE)
    except PermissionError:
        print('Could not save the spreadsheet (is it open in Excel?). What you typed was still used;')
        print('close the spreadsheet and run again to add new games to it.\n')

    filled = sum(1 for r in rows if r[2] or r[3])
    print(f'Spreadsheet: {len(rows)} games on the {SHEET_NAME} tab, {filled} with a genre or complexity.')
    for problem in problems:
        print(f'  Check: {problem}')
    print()
    return True


def inherit_from_base_game(games):
    """Expansions live on the shelf with their base game, so they share its
    genre, complexity and 18+ tag unless they have their own."""
    by_id = {g['id']: g for g in games}
    for game in games:
        base = by_id.get((game.get('baseGame') or {}).get('id'))
        if not base:
            continue
        for field in ('shelfLocation', 'ourComplexity', 'adultsOnly'):
            if not game.get(field):
                game[field] = base.get(field, CUSTOM_FIELDS[field])


def save_output(output):
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(output, f, indent=2, ensure_ascii=False)


def push_to_github(message):
    """Commit games.json and push it so the Square website updates."""
    print('Pushing update to GitHub...')
    try:
        subprocess.run(['git', 'add', 'games.json'], cwd=SCRIPT_DIR, check=True)
        unchanged = subprocess.run(['git', 'diff', '--cached', '--quiet', '--', 'games.json'], cwd=SCRIPT_DIR)
        if unchanged.returncode == 0:
            print('games.json has no changes, so there was nothing to push.\n')
            return
        subprocess.run(['git', 'commit', '-m', message, '--', 'games.json'], cwd=SCRIPT_DIR, check=True)
        subprocess.run(['git', 'push'], cwd=SCRIPT_DIR, check=True)
        print('Pushed to GitHub! Your Square website will update shortly.\n')
    except (subprocess.CalledProcessError, FileNotFoundError) as e:
        print(f'Git push failed: {e}')
        print('You may need to push manually: git add games.json && git commit -m "Update library" && git push\n')


def publish_sheet_changes(push):
    """--sheet-only: apply the spreadsheet to games.json without asking BGG."""
    print('=' * 50)
    print('Publishing shelf info from the spreadsheet')
    print('=' * 50 + '\n')
    with open(OUTPUT_FILE, 'r', encoding='utf-8') as f:
        output = json.load(f)
    games = output['games']
    for game in games:
        for field, default in CUSTOM_FIELDS.items():
            game.setdefault(field, default)
    if not sync_shelf_sheet(games):
        return
    inherit_from_base_game(games)
    save_output(output)
    if push:
        push_to_github('Update genre and complexity from the spreadsheet')


def main():
    push = '--no-push' not in sys.argv
    if '--sheet-only' in sys.argv:
        publish_sheet_changes(push)
        return

    print('=' * 50)
    print('BGG Library Updater')
    print('=' * 50)
    print(f'\nFetching collection for: {BGG_USERNAME}\n')

    # Load existing data to preserve custom fields
    existing_data = {}
    if os.path.exists(OUTPUT_FILE):
        try:
            with open(OUTPUT_FILE, 'r', encoding='utf-8') as f:
                existing = json.load(f)
                for game in existing.get('games', []):
                    existing_data[game['id']] = {
                        field: game.get(field, default) for field, default in CUSTOM_FIELDS.items()
                    }
            print(f'Found existing data for {len(existing_data)} games\n')
        except Exception:
            print('No existing data found, starting fresh\n')

    # Fetch collection from BGG
    print('Fetching collection from BGG...')
    collection_url = (
        f'https://boardgamegeek.com/xmlapi2/collection'
        f'?username={BGG_USERNAME}&own=1&stats=1'
    )
    collection_xml = fetch_url(collection_url)

    # Parse collection
    games = parse_collection(collection_xml)
    print(f'\nFound {len(games)} games in collection\n')

    if not games:
        print('No games found! Check that your BGG username is correct and your collection is public.')
        return

    # Fetch detailed info (descriptions, designers, categories, mechanics, etc.)
    print('Fetching game details (this may take a few minutes)...\n')
    fetch_game_details(games)

    # Merge with existing custom data
    for game in games:
        if game['id'] in existing_data:
            game.update(existing_data[game['id']])

    # Sort by name
    games.sort(key=lambda g: g['name'].lower())

    # Pull in what's been typed into the spreadsheet, and add new games to it
    sync_shelf_sheet(games)
    inherit_from_base_game(games)

    # Save to file
    save_output({
        'lastUpdated': datetime.now().isoformat(),
        'username': BGG_USERNAME,
        'totalGames': len(games),
        'games': games,
    })

    # Summary stats
    with_desc = sum(1 for g in games if g['description'])
    with_designers = sum(1 for g in games if g['designers'])
    with_honors = sum(1 for g in games if g['honors'])

    print('=' * 50)
    print(f'SUCCESS! Saved {len(games)} games to games.json')
    print(f'  - {with_desc} with descriptions')
    print(f'  - {with_designers} with designer info')
    print(f'  - {with_honors} with awards/honors')
    print('=' * 50)
    print('\nYour custom data (genre, complexity, small area, 18+) is preserved across updates.\n')

    # Auto-commit and push to GitHub so the Square website updates
    if push:
        push_to_github('Update game library data from BGG')


if __name__ == '__main__':
    main()
